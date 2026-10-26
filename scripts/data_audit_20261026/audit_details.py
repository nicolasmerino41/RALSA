from pathlib import Path
import csv,json,sys,re,hashlib,collections,math
import pandas as pd
import openpyxl
archive=Path(sys.argv[1]);repo=Path(sys.argv[2]);out=Path(sys.argv[3]);out.mkdir(parents=True,exist_ok=True)
names=['Quercus','Nahuel','Salix_Galpar','Gottin_HP','Gottin_PP','Garraf_HP','Garraf_PP','Garraf_PP2','Olot','Montseny']
inputs=repo/'outputs/support_prediction_zero/confirmed/inputs'
sites={ds:set(pd.read_csv(inputs/f'{ds}_occupancy.csv',dtype=str).site) for ds in names}
current={ds:pd.read_csv(repo/f'data/processed/{ds}_interactions.csv',dtype=str) for ds in names}
evidence={};extra=[]
def track(p):
 extra.append(dict(path=str(p),bytes=p.stat().st_size,sha256=hashlib.sha256(p.read_bytes()).hexdigest()))
def read(p,**kw):
 track(p);return pd.read_csv(p,**kw)
def triples(df):return set(df[['site','consumer','resource']].itertuples(index=False,name=None))
def num(x):
 try:return float(str(x).replace(',','.'))
 except:return float('nan')
def clean(x):return str(x).strip() if x is not None else ''
def rname(x):return re.sub(r'[^\w.]','.',clean(x))
for ds in names:
 track(inputs/f'{ds}_occupancy.csv');track(repo/f'data/processed/{ds}_interactions.csv')

# Raw Salix includes one row per rearing lot, with zeros retained for parasitoids.
p=archive/'Salix/raw-data/Salix_webs.csv';track(p)
raw=list(csv.reader(p.open(encoding='utf-8-sig')))
head=raw[10][:24];d=pd.DataFrame([r[:24] for r in raw[11:]],columns=head)
print('SALIX HEAD',head)
parids=raw[9][24:];entries=[]
for row in raw[11:]:
 for j,v in enumerate(row[24:]):
  if num(v)>0:entries.append((row[5],parids[j],row[14],num(v),row[1]))
rawint=pd.DataFrame(entries,columns=['site','consumer','resource','count','rearing_number'])
curr=triples(current['Salix_Galpar']);ri=triples(rawint)
sp=[]
for site,v in d.groupby('site',sort=False):
 coords=set(zip(v.Ndecdeg.map(num),v.Edecdeg.map(num)))
 dates=set(v['leg']);years=set(v['Year of coll'])
 sp.append(dict(site=site,in_current=site in sites['Salix_Galpar'],rearing_lots=len(v),dates=len(dates),years=len(years),first_year=min(years),last_year=max(years),coordinate_pairs=len(coords),latitude_min=min(x[0] for x in coords),latitude_max=max(x[0] for x in coords),longitude_min=min(x[1] for x in coords),longitude_max=max(x[1] for x in coords),has_parasitoid_record=site in set(rawint.site)))
sd=pd.DataFrame(sp);sd.to_csv(out/'salix_site_metadata_audit.csv',index=False)
sg=d.groupby(['site','Rgaller']).agg(lots=('Rearing number','size'),dates=('leg','nunique'),years=('Year of coll','nunique')).reset_index();sg.to_csv(out/'salix_site_host_replication.csv',index=False)
ev=dict(raw_lots=len(d),raw_sites=d.site.nunique(),current_sites=len(sites['Salix_Galpar']),raw_interacting_sites=rawint.site.nunique(),current_sites_missing_raw=sorted(sites['Salix_Galpar']-set(d.site)),raw_only_sites=sorted(set(d.site)-sites['Salix_Galpar']),raw_triples=len(ri),current_triples=len(curr),current_not_raw=len(curr-ri),raw_not_current=len(ri-curr),year_min=min(d['Year of coll']),year_max=max(d['Year of coll']),distinct_years=d['Year of coll'].nunique(),multiple_date_sites=int((sd.dates>1).sum()),current_multiple_date_sites=int(((sd.dates>1)&sd.in_current).sum()),current_multiple_year_sites=int(((sd.years>1)&sd.in_current).sum()),current_multiple_coordinate_sites=int(((sd.coordinate_pairs>1)&sd.in_current).sum()),site_host_multiple_dates=int((sg.dates>1).sum()),site_host_combinations=len(sg),count_max=float(rawint['count'].max()),nonpositive_ngalls=int((d['n galls'].map(num)<=0).sum()),coordinate_invalid_rows=int((~d.Ndecdeg.map(num).between(-90,90)|~d.Edecdeg.map(num).between(-180,180)).sum()))
evidence['Salix_Galpar']=ev
pd.DataFrame(sorted(ri-curr),columns=['site','consumer','resource']).to_csv(out/'salix_raw_only_interaction_triples.csv',index=False)

# Göttingen: preserve record-level fields; compare parasite slots separately.
hp=read(archive/'Gottin/raw-data/host_para_all_interactions.csv',sep=';',dtype=str,keep_default_na=False)
pp=read(archive/'Gottin/raw-data/plant_poll_all_interactions.csv',sep=';',dtype=str,keep_default_na=False)
gm=read(archive/'Gottin/metadata.csv',dtype=str,keep_default_na=False)
rows=[]
for slot in [1,2,3]:
 for i,r in hp.iterrows():
  if r[f'P{slot}.Species.Genus'] not in ['','NA'] and num(r[f'P{slot}.cells'])>0:
   rows.append(dict(source_record=i+1,site=r.Site,consumer=r[f'P{slot}.Species.Genus'],resource=r['Genus.Species'],count=num(r[f'P{slot}.cells']),slot=slot))
hrs=pd.DataFrame(rows);hcur=triples(current['Gottin_HP']);hknown=hrs[hrs.site.isin(sites['Gottin_HP'])]
for ds in ['Gottin_HP','Gottin_PP']:
 evidence[ds]={'metadata_matching_sites':len(sites[ds]&set(gm.site)), 'metadata_fields':list(gm.columns),'coordinates_present':False}
evidence['Gottin_HP'].update(dict(records=len(hp),records_missing_site=int((hp.Site=='').sum()),slots=[dict(slot=i,positive_identified_records=int((hrs.slot==i).sum()),triples=len(triples(hknown[hknown.slot==i])),new_current_triples=len(triples(hknown[hknown.slot==i])-hcur)) for i in [1,2,3]],total_new_triples=len(triples(hknown)-hcur),new_regional_pairs=len(set(zip(hknown.consumer,hknown.resource))-set(zip(current['Gottin_HP'].consumer,current['Gottin_HP'].resource))),nest_ids=sorted(hp.nest_ID.unique()),spec_id_unique=hp.spec_ID.nunique(),positive_cells=int((hp.cells.map(num)>0).sum()),unidentified_positive_P1=int(((hp['P1.Species.Genus']=='')&(hp['P1.cells'].map(num)>0)).sum())))
hknown[[(r.site,r.consumer,r.resource) not in hcur for r in hknown.itertuples()]].to_csv(out/'gottin_HP_additional_slot_records.csv',index=False)
hp.groupby('Site').agg(records=('spec_ID','size'),nest_ids=('nest_ID','nunique'),spec_ids=('spec_ID','nunique')).to_csv(out/'gottin_HP_replication.csv')
present=pp[~pp['P.Genus.Species'].isin(['','NA'])].rename(columns={'Site':'site','Genus.Species':'consumer','P.Genus.Species':'resource'})
pd.DataFrame(sorted(triples(present)-triples(current['Gottin_PP'])),columns=['site','consumer','resource']).to_csv(out/'gottin_PP_raw_only_triples.csv',index=False)
evidence['Gottin_PP'].update(dict(records=len(pp),known_plant_records=len(present),missing_plant_records=len(pp)-len(present),raw_insect_species=pp['Genus.Species'].nunique(),raw_known_plant_triples=len(triples(present)),current_triples=len(current['Gottin_PP']),raw_extra_triples=len(triples(present)-triples(current['Gottin_PP'])),transect_values=sorted(pp.Transect.unique()),section_values=sorted(pp.T_section.unique()),site_transect_section_groups=len(pp.groupby(['Site','Transect','T_section'])),dates_available=False))
pp.groupby('Site').agg(records=('Transect','size'),transects=('Transect','nunique'),section_codes=('T_section','nunique')).to_csv(out/'gottin_PP_replication.csv')

# Matrix contents, excluding taxon headers and workbook summaries.
matrix=[];qu=[]
for p in sorted((archive/'Quercus/raw-data').glob('*.xlsx')):
 track(p);site,year=re.match(r'Web(\d+) (\d+)',p.name).groups()
 wb=openpyxl.load_workbook(p,data_only=True);rows=list(wb['Sheet1'].values);head=rows[0]
 cells=[(clean(r[0]),clean(head[j]),num(r[j])) for r in rows[1:] if r[0] is not None for j in range(1,len(head)) if head[j] is not None]
 ints=[(site,rname(c),r,v) for r,c,v in cells if c.lower()!='unpar' and v>0]
 un=[v for r,c,v in cells if c.lower()=='unpar' and v>=0]
 qu.append(dict(site=site,year=year,positive_interaction_cells=len(ints),interaction_total=sum(v[3] for v in ints),unparasitised_total=sum(un),host_rows=sum(r[0] is not None for r in rows[1:]),parasitoid_columns=sum(x is not None and clean(x).lower()!='unpar' for x in head[1:]),maximum_interaction=max([x[3] for x in ints],default=0)))
 wb.close()
pd.DataFrame(qu).to_csv(out/'quercus_site_year_inventory.csv',index=False)
qm=read(archive/'Quercus/metadata.csv',dtype=str,keep_default_na=False);qm=qm[qm.Tree!=''];qm.to_csv(out/'quercus_coordinates_as_recorded.csv',columns=['Tree','East','North','Connectivity'],index=False)
evidence['Quercus']=dict(years=sorted(set(x['year'] for x in qu)),files=len(qu),sites_by_year={y:len({x['site'] for x in qu if x['year']==y}) for y in ['2006','2007']},coordinate_matches=len(set(qm.Tree)&sites['Quercus']),coordinate_system='not specified in held metadata',within_year_dates=False)
for ds,name in [('Garraf_HP','garraf-hp'),('Garraf_PP','garraf-pp'),('Garraf_PP2','garraf-pp-2'),('Olot','olot'),('Montseny','montseny')]:
 p=archive/f'Garraf-Montseny-Olot/raw-data/{name}.xlsx';track(p);wb=openpyxl.load_workbook(p,data_only=True)
 extra_sheets=set(wb.sheetnames)-sites[ds]
 for site in sorted(sites[ds]):
  ws=wb[site];rows=list(ws.values);head=rows[0]
  numeric=[num(r[j]) for r in rows[1:] if r[0] is not None for j in range(1,len(head)) if head[j] is not None and math.isfinite(num(r[j]))]
  matrix.append(dict(dataset=ds,site=site,numeric_cells=len(numeric),positive_cells=sum(x>0 for x in numeric),maximum=max(numeric,default=0),sum_values=sum(numeric),noninteger_cells=sum(x!=int(x) for x in numeric),values_gt1=sum(x>1 for x in numeric)))
 evidence[ds]=dict(site_sheets=len(sites[ds]&set(wb.sheetnames)),other_sheets=sorted(extra_sheets),hidden_sheets=[ws.title for ws in wb if ws.sheet_state!='visible'],hidden_columns={ws.title:[k for k,v in ws.column_dimensions.items() if v.hidden] for ws in wb if any(v.hidden for v in ws.column_dimensions.values())},notes=[])
 for sh in extra_sheets-{'metaweb'}:
  ws=wb[sh]
  evidence[ds]['notes'] += [dict(sheet=sh,row=i+1,text=' | '.join(str(x) for x in r if x is not None)) for i,r in enumerate(ws.values) if any(x is not None for x in r)][:210]
 if ds=='Garraf_PP':
  c=pd.DataFrame(list(wb['coordinates'].values)[1:],columns=['Plot','X','Y']).dropna(how='all');c['site']='plot'+c.Plot.astype(int).astype(str);c.to_csv(out/'garraf_PP_coordinates_as_recorded.csv',index=False)
  evidence[ds]['coordinate_matches']=len(set(c.site)&sites[ds]);evidence[ds]['coordinate_system']='X/Y numeric values consistent with longitude/latitude; datum undocumented'
 wb.close()
pd.DataFrame(matrix).to_csv(out/'matrix_quantitative_inventory.csv',index=False)
evidence['Nahuel']=dict(files=8,format='90 x 14 numeric matrices; no taxa/site metadata columns',site_mapping='vaz_ag=1,vaz_cl=2,vaz_ll=3,vaz_mh=4,vaz_mnh=5,vaz_qh=6,vaz_qnh=7,vaz_s=8')
for p in (archive/'csv').glob('df_*.csv'):track(p)
(out/'detailed_evidence.json').write_text(json.dumps(evidence,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
pd.DataFrame(extra).drop_duplicates('path').to_csv(out/'additional_input_versions.csv',index=False)
for ds,ev in evidence.items():print(ds,json.dumps({k:v for k,v in ev.items() if k!='notes'},ensure_ascii=True))
