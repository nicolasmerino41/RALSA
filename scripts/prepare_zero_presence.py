from pathlib import Path
import csv, hashlib, json, shutil, sys
import pandas as pd
import openpyxl

repo=Path(sys.argv[1])
archive=Path(sys.argv[2])
dest=Path(sys.argv[3]) if len(sys.argv)>3 else repo/'outputs/support_prediction_zero'
dest.mkdir(parents=True,exist_ok=True)
datasets=[p.name.removesuffix('_occupancy.csv') for p in (repo/'data/processed').glob('*_occupancy.csv')]
base={d:pd.read_csv(repo/f'data/processed/{d}_occupancy.csv',dtype=str) for d in datasets}
additions=[];sources={}
def track(path):
    sources[str(path)]=hashlib.sha256(path.read_bytes()).hexdigest()
def add(ds,site,species,guild,evidence,path,location,uncertain=False):
    site=str(site).strip();species=str(species).strip()
    if not site or site not in set(base[ds].site) or species in ['', 'None','NA','NaN']:return
    additions.append(dict(dataset=ds,site=site,species=species,trophic_level=guild,evidence=evidence,source=str(path),location=str(location),uncertain=uncertain))
def num(x):
    try:return float(x or 0)
    except (ValueError,TypeError):return 0
meta=archive/'Quercus/metadata.csv';track(meta)
trees=pd.read_csv(meta,dtype=str).Tree.iloc[:22]
for site in trees:
    path=archive/f'Quercus/raw-data/Web{site} 2006 Kaartinen.xlsx';track(path)
    wb=openpyxl.load_workbook(path,data_only=True,read_only=True)
    rows=list(wb['Sheet1'].values);head=rows[0]
    u=next(j for j,x in enumerate(head) if str(x).lower()=='unpar')
    for i,r in enumerate(rows[1:],2):
        if r[0] is not None and num(r[u])>0:add('Quercus',site,r[0],'resource','positive unparasitised host count',path,f'Sheet1 row {i}')
    wb.close()
for ds,file in [('Garraf_HP','garraf-hp'),('Olot','olot'),('Garraf_PP2','garraf-pp-2'),('Montseny','montseny')]:
    path=archive/f'Garraf-Montseny-Olot/raw-data/{file}.xlsx';track(path)
    wb=openpyxl.load_workbook(path,data_only=True,read_only=True)
    sheets=wb.sheetnames if ds=='Garraf_PP2' else wb.sheetnames[1:-1]
    for site in sheets:
        rows=list(wb[site].values);head=rows[0];cols=[j for j,x in enumerate(head) if j>0 and x is not None]
        rows=[r for r in rows[1:] if r[0] is not None]
        if ds=='Garraf_PP2':
            for j in cols:
                if all(num(r[j])==0 for r in rows):add(ds,site,head[j],'resource','named all-zero plant column; presence interpretation uncertain',path,f'{site} column {j+1}',True)
        else:
            for i,r in enumerate(rows,2):
                if all(num(r[j])==0 for j in cols):
                    species=('Plant-' if ds=='Montseny' else '')+str(r[0])
                    add(ds,site,species,'resource','INFO explicitly retains non-parasitised hosts' if ds!='Montseny' else 'named all-zero plant row; presence interpretation uncertain',path,f'{site} row {i}',ds=='Montseny')
    wb.close()
for ds,file in [('Gottin_HP','host_para_all_interactions'),('Gottin_PP','plant_poll_all_interactions')]:
    path=archive/f'Gottin/raw-data/{file}.csv';track(path)
    with path.open(encoding='utf-8-sig') as f:rows=list(csv.DictReader(f,delimiter=';'))
    for i,r in enumerate(rows,2):
        if ds=='Gottin_HP' and num(r['cells'])>0:add(ds,r['Site'],r['Genus.Species'],'resource','positive host cell count',path,i)
        if ds=='Gottin_PP' and r['P.Genus.Species'].strip() in ['','NA','NaN']:
            add(ds,r['Site'],r['Genus.Species'],'consumer','named insect with missing plant identity; sampling interpretation uncertain',path,i,True)
audit=pd.DataFrame(additions)
keys=['site','species','trophic_level']
audit['already_present']=[tuple(r[k] for k in keys) in set(map(tuple,base[r.dataset][keys].to_numpy())) for _,r in audit.iterrows()]
audit.to_csv(dest/'presence_record_provenance.csv',index=False)
pd.DataFrame([dict(path=k,sha256=v) for k,v in sources.items()]).to_csv(dest/'source_checksums.csv',index=False)
summary=[]
for scenario in ['confirmed','inclusive']:
    inputs=dest/scenario/'inputs';inputs.mkdir(parents=True,exist_ok=True)
    for ds in datasets:
        a=audit[(audit.dataset==ds)&((~audit.uncertain) if scenario=='confirmed' else True)]
        occ=pd.concat([base[ds],a[keys]]).drop_duplicates().sort_values(keys)
        oldco=pd.read_csv(repo/f'data/processed/{ds}_cooccurrences.csv',dtype=str)
        ints=pd.read_csv(repo/f'data/processed/{ds}_interactions.csv',dtype=str)
        co=occ[occ.trophic_level=='consumer'][['site','species']].rename(columns={'species':'consumer'}).merge(occ[occ.trophic_level=='resource'][['site','species']].rename(columns={'species':'resource'}),on='site')
        co=co[list(oldco.columns)].drop_duplicates().sort_values(list(oldco.columns))
        assert set(map(tuple,oldco.to_numpy()))<=set(map(tuple,co.to_numpy()))
        assert set(map(tuple,ints.to_numpy()))<=set(map(tuple,co.to_numpy()))
        occ.to_csv(inputs/f'{ds}_occupancy.csv',index=False)
        co.to_csv(inputs/f'{ds}_cooccurrences.csv',index=False)
        shutil.copyfile(repo/f'data/processed/{ds}_interactions.csv',inputs/f'{ds}_interactions.csv')
        oldpairs=len(oldco[['consumer','resource']].drop_duplicates());newpairs=len(co[['consumer','resource']].drop_duplicates())
        summary.append(dict(scenario=scenario,dataset=ds,added_presences=len(occ)-len(base[ds]),old_opportunities=len(oldco),new_opportunities=len(co),old_pairs=oldpairs,new_pairs=newpairs,old_p_emp=len(ints)/len(oldco),new_p_emp=len(ints)/len(co)))
pd.DataFrame(summary).to_csv(dest/'input_changes.csv',index=False)
print(pd.DataFrame(summary).to_string(index=False))
